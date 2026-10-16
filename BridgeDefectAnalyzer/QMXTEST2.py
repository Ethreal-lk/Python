import pandas as pd
from core.defect_rules import (
    DECK_COMPONENT_DEFECT_RULES,
    DECK_COMPONENT_ALIAS,
    LOWER_COMPONENT_DEFECT_RULES,
    LOWER_COMPONENT_ALIAS,
    LOWER_DEFECT_MAP,
    DECK_COMPONENT_TYPES,
)
from core.excel_handler import read_excel
from core.eport_generator import process_bridge_summary_to_new_file
from core.inspection_filler import parse_excel_if_formula
from openpyxl import load_workbook
import os
import sys
import re
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl import Workbook
sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")

# 4. 补全缺失的损坏类型  ← 加在这里
DECK_PART_DEFECT_TYPES = {
        component: list(categories.keys())
        for component, categories in DECK_COMPONENT_DEFECT_RULES.items()
    }
# 4. 补全缺失的损坏类型  ← 加在这里
PART_DEFECT_TYPES = {
        component: list(categories.keys())
        for component, categories in LOWER_COMPONENT_DEFECT_RULES.items()
    }

def traverse_bridge_data(df):
    """
    按 桥梁 → 孔跨 → 部位 → 部件 → 构件 → 损坏类型 层级遍历，
    打印层级结果，并返回 DataFrame。
    """
    records = []

    for bridge, df_bridge in df.groupby("桥梁名称"):
        print(f"桥梁：{bridge}")
        for span, df_span in df_bridge.groupby("孔跨"):
            print(f"  孔跨：{span}")
            for part, df_part in df_span.groupby("部位类型"):
                print(f"    部位：{part}")
                for comp, df_comp in df_part.groupby("部件类型"):
                    print(f"      部件：{comp}")
                    for member, df_member in df_comp.groupby("构件名称"):
                        print(f"        构件：{member}")
                        for defect, df_defect in df_member.groupby("损坏类型"):
                            max_score = df_defect["扣分分数"].max()
                            print(f"          损坏类型：{defect}，最大扣分：{max_score}")
                            records.append({
                                "桥梁名称": bridge,
                                "孔跨": span,
                                "部位类型": part,
                                "部件类型": comp,
                                "构件名称": member,
                                "损坏类型": defect,
                                "扣分分数": max_score,
                            })

    return pd.DataFrame(records)

def parse_template_rules(template_path, sheet_name):
    """
    从模板解析 (构件类型, 损坏类型) -> {扣分: 程度文本}
    """
    defect_rule_map = {}
    wb = load_workbook(template_path, data_only=False)
    ws = wb[sheet_name]

    print(f"正在解析模板: {template_path} | Sheet: {sheet_name}")

    current_component = None

    for row in ws.iter_rows(min_row=3, values_only=False):
        component_cell = row[0].value   # A列：构件类型
        defect_type = row[1].value      # B列：损坏类型
        formula = row[3].value          # D列：公式

        # 处理合并单元格：A列有值时更新当前构件
        if component_cell:
            current_component = str(component_cell).strip()

        if defect_type and formula and current_component:
            key = (current_component, str(defect_type).strip())
            rule_map, default_deg = parse_excel_if_formula(str(formula))
            if rule_map:
                defect_rule_map[key] = rule_map
                print(f"  ✅ ({current_component}, {defect_type}): {rule_map}")

    wb.close()
    print(f"共解析出 {len(defect_rule_map)} 条规则。")
    return defect_rule_map


def score_to_degree(score, rule_map):
    if not rule_map:
        return "无"
    if score in rule_map:
        return rule_map[score]
    closest = min(rule_map.keys(), key=lambda k: abs(k - score))
    return rule_map[closest]


def fill_missing_defects(df_result, part_defect_types, alias_map):
    records = []

    for (bridge, span, member), group in df_result.groupby(
        ["桥梁名称", "孔跨", "构件名称"]
    ):
        raw_part = group["部件类型"].iloc[0]
        std_part = alias_map.get(raw_part, raw_part)
        all_defects = part_defect_types.get(std_part, [])
        # 加在这里
        print(f"构件 {member}: 原始部件={raw_part}, 标准部件={std_part}, 损坏类型数={len(all_defects)}")
        existing_score = dict(zip(group["损坏类型"], group["扣分分数"]))
        existing_degree = dict(zip(group["损坏类型"], group["程度"]))

        for defect in all_defects:
            records.append({
                "桥梁名称": bridge,
                "孔跨": span,
                "构件名称": member,
                "部件类型": std_part,
                "损坏类型": defect,
                "扣分分数": existing_score.get(defect, 0),
                "程度": existing_degree.get(defect, "无"),
            })

    return pd.DataFrame(records)

# 1. 数据转换：把某构件的记录转成画表需要的字典
# ============================================================
def build_component_data(df_member, defect_order):
    """
    df_member: 某个构件的记录，列含 损坏类型、程度、扣分分数
    defect_order: 该部件下损坏类型的固定顺序
    返回：{损坏类型: {"degree": 程度, "score": 扣分}}
    """
    existing = dict(
        zip(df_member["损坏类型"], zip(df_member["程度"], df_member["扣分分数"]))
    )

    data = {}
    for dt in defect_order:
        degree, score = existing.get(dt, ("无", 0))
        data[dt] = {"degree": degree, "score": score}
    return data


# ============================================================
# 2. 画部件公共表头
# ============================================================
def draw_part_header(ws, start_row, defect_order):
    """
    画一个部件的公共表头，返回下一行行号。
    表头：A/B 列合并为斜线表头，C 列往后为损坏类型。
    """
    r = start_row

    # 左上角：损害类型 \n 结构位置
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ws.cell(row=r, column=1, value="损害类型\n结构位置")
    ws.cell(row=r, column=1).alignment = Alignment(
        horizontal="left", vertical="center", wrap_text=True
    )

    # 损坏类型列
    for idx, defect in enumerate(defect_order):
        col = 3 + idx
        cell = ws.cell(row=r, column=col, value=defect)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 样式
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    max_col = 2 + len(defect_order)

    for c in range(1, max_col + 1):
        cell = ws.cell(row=r, column=c)
        cell.border = border
        cell.font = Font(name="宋体", size=11, bold=True)

    # 左上角斜线
    diag = Border(
        left=thin, right=thin, top=thin, bottom=thin,
        diagonal=thin, diagonalDown=True,
    )
    ws.cell(row=r, column=1).border = diag
    ws.cell(row=r, column=2).border = diag
    ws.cell(row=r, column=1).alignment = Alignment(
        horizontal="left", vertical="center", wrap_text=True
    )

    # 行高列宽
    ws.row_dimensions[r].height = 40
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 16)
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 14)
    for idx in range(len(defect_order)):
        col_letter = get_column_letter(3 + idx)
        ws.column_dimensions[col_letter].width = max(
            ws.column_dimensions[col_letter].width or 0, 16
        )

    return r + 1


# ============================================================
# 3. 画单个构件的两行（程度、扣分值）
# ============================================================
def draw_component_rows(ws, start_row, component_name, component_data, defect_order):
    """
    画一个构件的两行，共用上面的表头，返回下一行行号。
    """
    r2 = start_row          # 程度行
    r3 = start_row + 1      # 扣分行

    # 构件名合并两行
    ws.merge_cells(start_row=r2, start_column=1, end_row=r3, end_column=1)
    ws.cell(row=r2, column=1, value=component_name)
    ws.cell(row=r2, column=1).alignment = Alignment(
        horizontal="center", vertical="center"
    )

    # 程度 / 扣分值
    ws.cell(row=r2, column=2, value="程度")
    ws.cell(row=r3, column=2, value="扣分值")

    # 数据
    for idx, defect in enumerate(defect_order):
        col = 3 + idx
        info = component_data.get(defect, {"degree": "无", "score": 0})
        ws.cell(row=r2, column=col, value=info["degree"])
        ws.cell(row=r3, column=col, value=info["score"])

    # 样式
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    max_col = 2 + len(defect_order)

    for r in (r2, r3):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.font = Font(name="宋体", size=11, bold=(c <= 2))

    ws.row_dimensions[r2].height = 25
    ws.row_dimensions[r3].height = 25

    return r3 + 1


# ============================================================
# 4. 调度：每个部件一个表头，构件共用表头往下排
# ============================================================
def draw_lower_structure(ws, df_lower):
    """
    在 ws 里画下部结构评分表。
    df_lower: 单座桥的下部结构数据，列含 构件名称、部件类型、损坏类型、扣分分数、程度
    返回：下一个表格的起始行号
    """
    current_row = 1

    for part, df_part in df_lower.groupby("部件类型"):
        std_part = part
        defect_order = PART_DEFECT_TYPES.get(std_part, [])
        if not defect_order:
            print(f"  ⚠️ 未找到 {std_part} 的损坏类型列表，跳过")
            continue

        # 1. 画部件公共表头
        current_row = draw_part_header(ws, current_row, defect_order)

        # 2. 每个构件画两行，共用表头
        for member, df_member in df_part.groupby("构件名称"):
            component_data = build_component_data(df_member, defect_order)
            current_row = draw_component_rows(
                ws, current_row, member, component_data, defect_order
            )

        current_row += 2   # 部件之间空 2 行

    return current_row

def draw_deck_summary(ws, df_deck):
    current_row = 1

    for part in DECK_COMPONENT_TYPES:          # 遍历所有部件，不是数据里的
        defect_order = DECK_PART_DEFECT_TYPES.get(part, [])
        if not defect_order:
            continue

        df_part = df_deck[df_deck["部件类型"] == part]
        score_map = dict(zip(df_part["损坏类型"], df_part["扣分分数"])) if not df_part.empty else {}
        degree_map = dict(zip(df_part["损坏类型"], df_part["程度"])) if not df_part.empty else {}

        r1 = current_row

        # 表头
        ws.merge_cells(start_row=r1, start_column=1, end_row=r1, end_column=2)
        ws.cell(row=r1, column=1, value="损害类型\n结构位置")
        ws.cell(row=r1, column=1).alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True
        )

        for idx, defect in enumerate(defect_order):
            col = 3 + idx
            cell = ws.cell(row=r1, column=col, value=defect)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # 部件名合并两行
        ws.merge_cells(start_row=r1 + 1, start_column=1, end_row=r1 + 2, end_column=1)
        ws.cell(row=r1 + 1, column=1, value=part)
        ws.cell(row=r1 + 1, column=1).alignment = Alignment(
            horizontal="center", vertical="center"
        )

        ws.cell(row=r1 + 1, column=2, value="程度")
        ws.cell(row=r1 + 2, column=2, value="扣分值")

        for idx, defect in enumerate(defect_order):
            col = 3 + idx
            score = score_map.get(defect, 0)
            degree = degree_map.get(defect, "无")
            ws.cell(row=r1 + 1, column=col, value=degree)
            ws.cell(row=r1 + 2, column=col, value=score)

        # 样式
        thin = Side(style="thin")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        max_col = 2 + len(defect_order)

        for r in range(r1, r1 + 3):
            for c in range(1, max_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = border
                if not (r == r1 and c == 1):
                    cell.alignment = Alignment(
                        horizontal="center", vertical="center", wrap_text=True
                    )
                cell.font = Font(name="宋体", size=11, bold=(r == r1 or c <= 2))

        diag = Border(
            left=thin, right=thin, top=thin, bottom=thin,
            diagonal=thin, diagonalDown=True,
        )
        ws.cell(row=r1, column=1).border = diag
        ws.cell(row=r1, column=2).border = diag

        ws.row_dimensions[r1].height = 40
        ws.row_dimensions[r1 + 1].height = 25
        ws.row_dimensions[r1 + 2].height = 25

        current_row = r1 + 4

    return current_row



def process_deck_sheet(wb, df_bridge, rule_file_path, safe_bridge_name):
    """
    处理桥面系 sheet，写入 wb。
    """
    ws_deck = wb.active
    ws_deck.title = "桥面系"

    df_deck = df_bridge[df_bridge["部位类型"] == "桥面系"]
    if df_deck.empty:
        ws_deck.cell(row=1, column=1, value="无桥面系数据")
        print(f"  ⚠️ [{safe_bridge_name}] 无桥面系数据。")
        return

    try:
        df_deck_result = traverse_bridge_data(df_deck)
        df_deck_rule_map = parse_template_rules(rule_file_path, "桥面系评分")

        df_deck_result["程度"] = df_deck_result.apply(
            lambda r: score_to_degree(
                r["扣分分数"],
                df_deck_rule_map.get(
                    (
                        DECK_COMPONENT_ALIAS.get(r["部件类型"], r["部件类型"]),
                        r["损坏类型"],
                    ),
                    {}
                )
            ),
            axis=1,
        )
        df_deck_result = fill_missing_defects(
            df_deck_result, DECK_PART_DEFECT_TYPES, DECK_COMPONENT_ALIAS
        )

        df_deck_summary = (
            df_deck_result
            .sort_values("扣分分数", ascending=False)
            .groupby(["桥梁名称", "部件类型", "损坏类型"], as_index=False)
            .first()
        )

        draw_deck_summary(ws_deck, df_deck_summary)
        print(f"  ✅ 桥面系处理成功。")
    except Exception as e:
        print(f"  ❌ 桥面系处理失败: {e}")
        ws_deck.cell(row=1, column=1, value=f"桥面系处理失败: {e}")
    
if __name__ == "__main__":
    # 1. 核心路径配置
    raw_defect_excel_path = r"E:\Python\BridgeDefectAnalyzer\input\病害列表导出2026-10-09_modify.xlsx"
    rule_file_path = r"E:\Python\BridgeDefectAnalyzer\input\城市桥梁评分计算.xlsx"
    output_dir = r"E:\Python\BridgeDefectAnalyzer\output"

    # 确保输出目录存在
    os.makedirs(output_dir, exist_ok=True)

    print(f"1. 正在读取全量病害总表: {raw_defect_excel_path}")
    df_all = read_excel(raw_defect_excel_path, skip_hidden_rows=True, skip_hidden_cols=True)
    # 提取桥面系数据
    df_deck_all = df_all[df_all["部位类型"] == "桥面系"].copy()
    print(f"桥面系记录数: {len(df_deck_all)}")
    print(df_deck_all)
    df_deck_result = traverse_bridge_data(df_deck_all)
    df_deck_rule_map = parse_template_rules(rule_file_path, "桥面系评分")
    df_deck_result["程度"] = df_deck_result.apply(
        lambda r: score_to_degree(
            r["扣分分数"],
            df_deck_rule_map.get(
                (
                    DECK_COMPONENT_ALIAS.get(r["部件类型"], r["部件类型"]),
                    r["损坏类型"],
                ),
                {}
            )
        ),
        axis=1,
    )
    df_deck_result = fill_missing_defects(df_deck_result, DECK_PART_DEFECT_TYPES, DECK_COMPONENT_ALIAS)
    print(df_deck_result)
    # 按桥 + 部件 + 损坏类型取最大扣分那行
    df_deck_summary = (
        df_deck_result
        .sort_values("扣分分数", ascending=False)
        .groupby(["桥梁名称", "部件类型", "损坏类型"], as_index=False)
        .first()
    )
    print(df_deck_summary)
    # ============================================================
    # 按桥画桥面系表  ← 加在这里
    # ============================================================
    for bridge, df_bridge in df_deck_summary.groupby("桥梁名称"):
        wb = Workbook()
        ws_deck = wb.active
        ws_deck.title = "桥面系"
        draw_deck_summary(ws_deck, df_bridge)

        out_path = os.path.join(output_dir, f"{bridge}_桥面系评分表.xlsx")
        wb.save(out_path)
        print(f"✅ 已生成: {out_path}")

