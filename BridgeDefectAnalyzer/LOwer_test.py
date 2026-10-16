import pandas as pd
from core.defect_rules import (
    DECK_COMPONENT_DEFECT_RULES,
    LOWER_COMPONENT_DEFECT_RULES,
    LOWER_COMPONENT_ALIAS,
    LOWER_DEFECT_MAP,
)
from core.excel_handler import read_excel
from core.eport_generator import process_bridge_summary_to_new_file
from core.inspection_filler import parse_excel_if_formula
from openpyxl import load_workbook
import os
import sys
import re
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl import Workbook
sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")

# 4. 补全缺失的损坏类型  ← 加在这里
PART_DEFECT_TYPES = {
        component: list(categories.keys())
        for component, categories in LOWER_COMPONENT_DEFECT_RULES.items()
    }
def traverse_bridge_data(df):
    """
    按 桥梁 → 孔跨 → 部位 → 部件 → 构件 → 损坏类型 层级遍历，
    打印层级结果，并返回 DataFrame。
    """
    records = []

    for bridge, df_bridge in df.groupby("桥梁名称"):
        print(f"桥梁：{bridge}")
        for span, df_span in df_bridge.groupby("孔跨"):
            print(f"  孔跨：{span}")
            for part, df_part in df_span.groupby("部位类型"):
                print(f"    部位：{part}")
                for comp, df_comp in df_part.groupby("部件类型"):
                    print(f"      部件：{comp}")
                    for member, df_member in df_comp.groupby("构件名称"):
                        print(f"        构件：{member}")
                        for defect, df_defect in df_member.groupby("损坏类型"):
                            max_score = df_defect["扣分分数"].max()
                            print(f"          损坏类型：{defect}，最大扣分：{max_score}")
                            records.append({
                                "桥梁名称": bridge,
                                "孔跨": span,
                                "部位类型": part,
                                "部件类型": comp,
                                "构件名称": member,
                                "损坏类型": defect,
                                "扣分分数": max_score,
                            })

    return pd.DataFrame(records)

def parse_template_rules(template_path, sheet_name):
    """
    从模板解析 (构件类型, 损坏类型) -> {扣分: 程度文本}
    """
    defect_rule_map = {}
    wb = load_workbook(template_path, data_only=False)
    ws = wb[sheet_name]

    print(f"正在解析模板: {template_path} | Sheet: {sheet_name}")

    current_component = None

    for row in ws.iter_rows(min_row=3, values_only=False):
        component_cell = row[0].value   # A列：构件类型
        defect_type = row[1].value      # B列：损坏类型
        formula = row[3].value          # D列：公式

        # 处理合并单元格：A列有值时更新当前构件
        if component_cell:
            current_component = str(component_cell).strip()

        if defect_type and formula and current_component:
            key = (current_component, str(defect_type).strip())
            rule_map, default_deg = parse_excel_if_formula(str(formula))
            if rule_map:
                defect_rule_map[key] = rule_map
                print(f"  ✅ ({current_component}, {defect_type}): {rule_map}")

    wb.close()
    print(f"共解析出 {len(defect_rule_map)} 条规则。")
    return defect_rule_map


def score_to_degree(score, rule_map):
    if not rule_map:
        return "无"
    if score in rule_map:
        return rule_map[score]
    closest = min(rule_map.keys(), key=lambda k: abs(k - score))
    return rule_map[closest]


def fill_missing_defects(df_result, part_defect_types, alias_map):
    records = []

    for (bridge, span, member), group in df_result.groupby(
        ["桥梁名称", "孔跨", "构件名称"]
    ):
        raw_part = group["部件类型"].iloc[0]
        std_part = alias_map.get(raw_part, raw_part)
        all_defects = part_defect_types.get(std_part, [])

        existing_score = dict(zip(group["损坏类型"], group["扣分分数"]))
        existing_degree = dict(zip(group["损坏类型"], group["程度"]))

        for defect in all_defects:
            records.append({
                "桥梁名称": bridge,
                "孔跨": span,
                "构件名称": member,
                "部件类型": std_part,
                "损坏类型": defect,
                "扣分分数": existing_score.get(defect, 0),
                "程度": existing_degree.get(defect, "无"),
            })

    return pd.DataFrame(records)

# 1. 数据转换：把某构件的记录转成画表需要的字典
# ============================================================
def build_component_data(df_member, defect_order):
    """
    df_member: 某个构件的记录，列含 损坏类型、程度、扣分分数
    defect_order: 该部件下损坏类型的固定顺序
    返回：{损坏类型: {"degree": 程度, "score": 扣分}}
    """
    existing = dict(
        zip(df_member["损坏类型"], zip(df_member["程度"], df_member["扣分分数"]))
    )

    data = {}
    for dt in defect_order:
        degree, score = existing.get(dt, ("无", 0))
        data[dt] = {"degree": degree, "score": score}
    return data


# ============================================================
# 2. 画部件公共表头
# ============================================================
def draw_part_header(ws, start_row, defect_order):
    """
    画一个部件的公共表头，返回下一行行号。
    表头：A/B 列合并为斜线表头，C 列往后为损坏类型。
    """
    r = start_row

    # 左上角：损害类型 \n 结构位置
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ws.cell(row=r, column=1, value="损害类型\n结构位置")
    ws.cell(row=r, column=1).alignment = Alignment(
        horizontal="left", vertical="center", wrap_text=True
    )

    # 损坏类型列
    for idx, defect in enumerate(defect_order):
        col = 3 + idx
        cell = ws.cell(row=r, column=col, value=defect)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 样式
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    max_col = 2 + len(defect_order)

    for c in range(1, max_col + 1):
        cell = ws.cell(row=r, column=c)
        cell.border = border
        cell.font = Font(name="宋体", size=11, bold=True)

    # 左上角斜线
    diag = Border(
        left=thin, right=thin, top=thin, bottom=thin,
        diagonal=thin, diagonalDown=True,
    )
    ws.cell(row=r, column=1).border = diag
    ws.cell(row=r, column=2).border = diag
    ws.cell(row=r, column=1).alignment = Alignment(
        horizontal="left", vertical="center", wrap_text=True
    )

    # 行高列宽
    ws.row_dimensions[r].height = 40
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 16)
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 14)
    for idx in range(len(defect_order)):
        col_letter = get_column_letter(3 + idx)
        ws.column_dimensions[col_letter].width = max(
            ws.column_dimensions[col_letter].width or 0, 16
        )

    return r + 1


# ============================================================
# 3. 画单个构件的两行（程度、扣分值）
# ============================================================
def draw_component_rows(ws, start_row, component_name, component_data, defect_order):
    """
    画一个构件的两行，共用上面的表头，返回下一行行号。
    """
    r2 = start_row          # 程度行
    r3 = start_row + 1      # 扣分行

    # 构件名合并两行
    ws.merge_cells(start_row=r2, start_column=1, end_row=r3, end_column=1)
    ws.cell(row=r2, column=1, value=component_name)
    ws.cell(row=r2, column=1).alignment = Alignment(
        horizontal="center", vertical="center"
    )

    # 程度 / 扣分值
    ws.cell(row=r2, column=2, value="程度")
    ws.cell(row=r3, column=2, value="扣分值")

    # 数据
    for idx, defect in enumerate(defect_order):
        col = 3 + idx
        info = component_data.get(defect, {"degree": "无", "score": 0})
        ws.cell(row=r2, column=col, value=info["degree"])
        ws.cell(row=r3, column=col, value=info["score"])

    # 样式
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    max_col = 2 + len(defect_order)

    for r in (r2, r3):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.font = Font(name="宋体", size=11, bold=(c <= 2))

    ws.row_dimensions[r2].height = 25
    ws.row_dimensions[r3].height = 25

    return r3 + 1


# ============================================================
# 4. 调度：每个部件一个表头，构件共用表头往下排
# ============================================================
def draw_lower_structure(ws, df_lower):
    """
    在 ws 里画下部结构评分表。
    df_lower: 单座桥的下部结构数据，列含 构件名称、部件类型、损坏类型、扣分分数、程度
    返回：下一个表格的起始行号
    """
    current_row = 1

    for part, df_part in df_lower.groupby("部件类型"):
        std_part = part
        defect_order = PART_DEFECT_TYPES.get(std_part, [])
        if not defect_order:
            print(f"  ⚠️ 未找到 {std_part} 的损坏类型列表，跳过")
            continue

        # 1. 画部件公共表头
        current_row = draw_part_header(ws, current_row, defect_order)

        # 2. 每个构件画两行，共用表头
        for member, df_member in df_part.groupby("构件名称"):
            component_data = build_component_data(df_member, defect_order)
            current_row = draw_component_rows(
                ws, current_row, member, component_data, defect_order
            )

        current_row += 2   # 部件之间空 2 行

    return current_row
def process_lower_sheet(wb, df_bridge, rule_file_path, safe_bridge_name):
    ws_lower = wb.create_sheet("下部结构")

    df_lower = df_bridge[df_bridge["部位类型"] == "下部结构"]
    print(f"  [调试] df_lower 行数: {len(df_lower)}")

    if df_lower.empty:
        ws_lower.cell(row=1, column=1, value="无下部结构数据")
        print(f"  ⚠️ [{safe_bridge_name}] 无下部结构数据。")
        return

    try:
        df_lower_result = traverse_bridge_data(df_lower)
        print(f"  [调试] traverse 后行数: {len(df_lower_result)}")

        if df_lower_result.empty:
            ws_lower.cell(row=1, column=1, value="下部结构无有效记录")
            print(f"  ⚠️ [{safe_bridge_name}] 下部结构无有效记录。")
            return

        df_lower_rule_map = parse_template_rules(rule_file_path, "下部结构评分")

        df_lower_result["程度"] = df_lower_result.apply(
            lambda r: score_to_degree(
                r["扣分分数"],
                df_lower_rule_map.get(
                    (LOWER_COMPONENT_ALIAS.get(r["部件类型"], r["部件类型"]), r["损坏类型"]),
                    {}
                )
            ),
            axis=1,
        )
        df_lower_result = fill_missing_defects(df_lower_result, PART_DEFECT_TYPES, LOWER_COMPONENT_ALIAS)

        df_lower_summary = (
            df_lower_result
            .sort_values("扣分分数", ascending=False)
            .groupby(["桥梁名称", "部件类型", "损坏类型"], as_index=False)
            .first()
        )

        draw_lower_structure(ws_lower, df_lower_summary)
        print(f"  ✅ 下部结构处理成功。")
    except Exception as e:
        print(f"  ❌ 下部结构处理失败: {e}")
        ws_lower.cell(row=1, column=1, value=f"下部结构处理失败: {e}")

# if __name__ == "__main__":
#     # 1. 核心路径配置
#     raw_defect_excel_path = r"E:\Python\BridgeDefectAnalyzer\input\病害列表导出2026-10-09_modify.xlsx"
#     rule_file_path = r"E:\Python\BridgeDefectAnalyzer\input\城市桥梁评分计算.xlsx"
#     output_dir = r"E:\Python\BridgeDefectAnalyzer\output"

#     # 确保输出目录存在
#     os.makedirs(output_dir, exist_ok=True)

#     print(f"1. 正在读取全量病害总表: {raw_defect_excel_path}")
#     df_all = read_excel(raw_defect_excel_path, skip_hidden_rows=True, skip_hidden_cols=True)
#     df_result = traverse_bridge_data(df_all)
#     defect_rule_map =parse_template_rules(rule_file_path, "下部结构评分")
#     df_result["程度"] = df_result.apply(
#     lambda r: score_to_degree(
#         r["扣分分数"],
#         defect_rule_map.get(
#             (LOWER_COMPONENT_ALIAS.get(r["部件类型"], r["部件类型"]), r["损坏类型"]),
#             {}
#         )
#     ),
#     axis=1,
#     )


#     df_result = fill_missing_defects(df_result, PART_DEFECT_TYPES, LOWER_COMPONENT_ALIAS)

#     for bridge, df_bridge in df_result.groupby("桥梁名称"):
#         wb = Workbook()

#         # Sheet 1: 桥面系（占位）
#         ws_deck = wb.active
#         ws_deck.title = "桥面系"
#         ws_deck.cell(row=1, column=1, value="桥面系评分表（待生成）")

#         # Sheet 2: 上部结构（占位）
#         ws_upper = wb.create_sheet("上部结构")
#         ws_upper.cell(row=1, column=1, value="上部结构评分表（待生成）")

#         # Sheet 3: 下部结构
#         ws_lower = wb.create_sheet("下部结构")
#         draw_lower_structure(ws_lower, df_bridge)

#         output_path = os.path.join(output_dir, f"{bridge}_评分表.xlsx")
#         wb.save(output_path)
#         print(f"✅ 已生成: {output_path}")


