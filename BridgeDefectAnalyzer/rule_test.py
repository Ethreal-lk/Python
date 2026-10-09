import re

from openpyxl import load_workbook

# 规则文件路径
rule_file_path = r"E:\Python\BridgeDefectAnalyzer\input\城市桥梁评分计算.xlsx"

def parse_excel_if_formula(formula_str):
    """
    解析 Excel 中的嵌套 IF 公式，提取扣分值与损坏程度的映射关系。
    例：'=IF(C10="无",0,IF(C10="<3%",5,IF(C10="3%~10%",15,IF(C10=">10%",40))))'
    返回映射字典: {0: '无', 5: '<3%', 15: '3%~10%', 40: '>10%'}
    """
    if not isinstance(formula_str, str) or not formula_str.startswith("="):
        return {}

    # 匹配类似于 Cxx="文本",数值 的正则模式
    pattern = r'C\d+\s*=\s*"([^"]+)"\s*,\s*(\d+)'
    matches = re.findall(pattern, formula_str)

    score_to_degree = {}
    for degree_text, score in matches:
        score_to_degree[int(score)] = degree_text

    return score_to_degree


def reverse_fill_and_save(excel_path, target_deductions, sheet_name="桥面系评分"):
    """
    根据给定的扣分目标，反推损坏程度，并填入 Excel 文件中保存
    :param excel_path: Excel 文件的绝对路径
    :param target_deductions: 目标扣分字典，如 {"缝内沉积物阻塞": 5, "丢残缺": 30}
    :param sheet_name: 目标工作表名称
    """
    print("1. 正在读取 Excel 规则和公式定义...")
    # 第一遍加载：data_only=False，用于解析 D 列中的 IF 公式
    wb_read = load_workbook(excel_path, data_only=False)
    ws_read = wb_read[sheet_name]

    defect_rules = {}  # 存放病害规则字典
    row_mapping = {}   # 存放病害对应的行号

    for row_idx, row in enumerate(ws_read.iter_rows(min_row=10, values_only=False), start=10):
        defect_type = row[1].value  # B列：损坏类型
        formula = row[3].value      # D列：单项扣分公式

        if defect_type and formula:
            rule_map = parse_excel_if_formula(str(formula))
            if rule_map:
                defect_rules[defect_type] = rule_map
                row_mapping[defect_type] = row_idx

    wb_read.close()
    print(f"   └─ 成功加载 {len(defect_rules)} 条病害扣分解析规则。")

    print("\n2. 开始反推损坏程度并写入 C 列...")
    # 第二遍加载：准备修改并保存文件
    wb_write = load_workbook(excel_path)
    ws_write = wb_write[sheet_name]

    updated_count = 0

    for defect_type, target_score in target_deductions.items():
        if defect_type not in defect_rules:
            print(f"⚠️  [跳过] 未在 Excel 中找到损坏类型【{defect_type}】的计算规则")
            continue

        rule_map = defect_rules[defect_type]
        row_num = row_mapping[defect_type]

        # 准确匹配分值，如果输入分值不在规则中，则匹配离它最近的分值
        if target_score in rule_map:
            degree_text = rule_map[target_score]
        else:
            closest_score = min(rule_map.keys(), key=lambda k: abs(k - target_score))
            degree_text = rule_map[closest_score]
            print(f"💡 [自动修正] 病害【{defect_type}】目标扣分 {target_score} 无精准对应，自动匹配离它最近的规则分值 {closest_score} -> '{degree_text}'")

        # 写入 C 列（第 3 列：损坏程度）
        ws_write.cell(row=row_num, column=3, value=degree_text)
        updated_count += 1
        print(f"✅ [填入成功] 行号 {row_num:<2} | 损坏类型: {defect_type:<14} | 扣分: {target_score:<2} -> 损坏程度: '{degree_text}'")

    # 保存更改
    wb_write.save(excel_path)
    wb_write.close()
    print(f"\n🎉 成功修改并保存了 {updated_count} 项数据！")
    print(f"📁 文件保存路径: {excel_path}")
    print("👉 现在可以使用 Excel / WPS 打开该文件，里面的公式会自动计算最新结果，或手动点击运行宏。")


# ==================== 运行测试 ====================


if __name__ == "__main__":
    # 模拟上游传过来的目标扣分数据
    test_target_deductions = {
        "网裂或龟裂": 5,      # 应该自动反推填充: "少量"
        "波浪及车辙": 15,          # 修正匹配或精准匹配 -> "5%~10%"
        "坑槽": 25,              # 应该自动反推填充: "2个"
        "桥面贯通横缝": 5,        # 应该自动反推填充: "3%~10%"
    }

    reverse_fill_and_save(rule_file_path, test_target_deductions)
