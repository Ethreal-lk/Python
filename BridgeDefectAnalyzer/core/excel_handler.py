import os 
import pandas as pd


import os
import pandas as pd
from openpyxl import load_workbook


def read_excel(file_path, sheet_name=0, skip_hidden_rows=False, skip_hidden_cols=False):
    """读取指定 Sheet，可跳过隐藏行/列"""

    if not os.path.exists(file_path):
        raise FileNotFoundError(
            "Excel 文件不存在: {}".format(file_path)
        )

    # 不需要跳过隐藏行、列，直接使用 pandas
    if not skip_hidden_rows and not skip_hidden_cols:
        return pd.read_excel(file_path, sheet_name=sheet_name)

    wb = load_workbook(file_path, data_only=True)

    try:
        ws = (
            wb.worksheets[sheet_name]
            if isinstance(sheet_name, int)
            else wb[sheet_name]
        )

        # 默认保留所有列
        hidden_cols = set()

        # 只在需要跳过隐藏列时，查找隐藏列
        if skip_hidden_cols:
            for col_letter, col_dim in ws.column_dimensions.items():
                if col_dim.hidden:
                    start = col_dim.min or column_index_from_string(col_letter)
                    end = col_dim.max or start
                    hidden_cols.update(range(start, end + 1))

        # 根据隐藏列索引，筛选需要保留的列
        visible_cols = [
            col_idx
            for col_idx in range(1, ws.max_column + 1)
            if col_idx not in hidden_cols
        ]

        # 读取表头
        headers = [
            ws.cell(row=1, column=col_idx).value
            for col_idx in visible_cols
        ]

        # 读取数据
        data = []

        for row_idx in range(2, ws.max_row + 1):

            # 按需跳过隐藏行
            if (
                skip_hidden_rows
                and ws.row_dimensions[row_idx].hidden
            ):
                continue

            row_data = [
                ws.cell(row=row_idx, column=col_idx).value
                for col_idx in visible_cols
            ]

            data.append(row_data)

        return pd.DataFrame(data, columns=headers)

    finally:
        wb.close()

def get_sheet_names(file_path):
    """获取 Sheet 名称"""
    if not os.path.exists(file_path):
        raise FileNotFoundError("Excel 文件不存在: {}".format(file_path))

    excel_file = pd.ExcelFile(file_path)
    return excel_file.sheet_names

def write_excel(data, file_path, sheet_name="Sheet1"):
    """写入单个 Sheet"""
    data.to_excel(
        file_path,
        sheet_name=sheet_name,
        index=False
    )

def write_sheets(data_dict, file_path):
    """写入多个 Sheet"""

    with pd.ExcelWriter(
        file_path,
        engine="openpyxl"
    ) as writer:

        for sheet_name, data in data_dict.items():
            data.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False
            )

def append_sheet(data, file_path, sheet_name):
    """追加 Sheet"""

    with pd.ExcelWriter(
        file_path,
        engine="openpyxl",
        mode="a",
        if_sheet_exists="replace"
    ) as writer:

        data.to_excel(
            writer,
            sheet_name=sheet_name,
            index=False
        )


def check_excel(file_path):
    """检查 Excel 是否可以正常读取"""

    try:
        get_sheet_names(file_path)
        return True
    except Exception:
        return False