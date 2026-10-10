import os
import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import pythoncom
import win32com.client




def force_excel_recalculate_and_save(file_path):
    """后台调用本地 Excel，强制重新计算公式并保存缓存"""
    abs_path = os.path.abspath(file_path)
    print(f"2. 正在通过后台 Excel 引擎强制计算并缓存公式: {abs_path}")

    pythoncom.CoInitialize()

    excel = None
    wb = None

    try:
        # 创建独立的 Excel 实例
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False

        # 注意：删除 excel.Calculation = -4105
        # 直接打开工作簿
        wb = excel.Workbooks.Open(
            abs_path,
            UpdateLinks=0,
            ReadOnly=False
        )

        # 强制完整重算并重建公式依赖关系
        print("   正在执行完整公式重算...")
        excel.CalculateFullRebuild()

        # 保存计算结果和公式缓存
        wb.Save()

        print("   ✅ Excel 公式计算与缓存保存成功！")

    except Exception as e:
        print(f"   ❌ 后台调用 Excel 失败: {e}")
        raise

    finally:
        if wb is not None:
            try:
                wb.Close(SaveChanges=False)
            except Exception as e:
                print(f"   ⚠️ 关闭工作簿失败: {e}")

        if excel is not None:
            try:
                excel.Quit()
            except Exception as e:
                print(f"   ⚠️ 退出 Excel 失败: {e}")

        wb = None
        excel = None

        pythoncom.CoUninitialize()


def draw_single_element_summary(ws, start_row, element_name, element_data):
  """核心画表函数：在新 Sheet 的特定起始行，绘制单个要素的扣分汇总表"""
  damage_types = list(element_data.keys())
  if not damage_types:
    return start_row

  r1 = start_row
  r2 = start_row + 1
  r3 = start_row + 2

  # 1. 写入主表头 (第 1 行)
  ws.merge_cells(
      start_row=r1, start_column=1, end_row=r1, end_column=2
  )
  ws.cell(
      row=r1,
      column=1,
      value="                   损害类型\n结构位置                   ",
  )
  ws.cell(row=r1, column=1).alignment = Alignment(
      horizontal="left", vertical="center", wrap_text=True
  )

  for idx, dt in enumerate(damage_types):
    col_idx = 3 + idx
    cell = ws.cell(row=r1, column=col_idx, value=dt)
    cell.alignment = Alignment(horizontal="center", vertical="center")

  # 2. 写入数据行（A列合并跨两行）
  ws.merge_cells(
      start_row=r2, start_column=1, end_row=r3, end_column=1
  )
  ws.cell(row=r2, column=1, value=element_name)

  ws.cell(row=r2, column=2, value="程度")
  ws.cell(row=r3, column=2, value="扣分值")

  for idx, dt in enumerate(damage_types):
    col_idx = 3 + idx
    info = element_data.get(dt, {"degree": "无", "score": 0})
    ws.cell(row=r2, column=col_idx, value=info["degree"])
    ws.cell(row=r3, column=col_idx, value=info["score"])

  max_col = 2 + len(damage_types)

  # 3. 样式精修（边框、斜线、居中、字体）
  thin_side = Side(style="thin", color="000000")
  normal_border = Border(
      left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
  )
  diagonal_border = Border(
      left=thin_side,
      right=thin_side,
      top=thin_side,
      bottom=thin_side,
      diagonal=thin_side,
      diagonalDown=True,
  )

  font_normal = Font(name="宋体", size=11, bold=False)
  font_bold = Font(name="宋体", size=11, bold=True)

  for r in range(r1, r3 + 1):
    for c in range(1, max_col + 1):
      cell = ws.cell(row=r, column=c)
      if not (r == r1 and c in [1, 2]):
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
      is_header = (r == r1) or (c <= 2) or (c == 1 and r in [r2, r3])
      cell.font = font_bold if is_header else font_normal
      cell.border = normal_border

  # 主表头左上角加上斜线边框
  ws.cell(row=r1, column=1).border = diagonal_border
  ws.cell(row=r1, column=2).border = diagonal_border
  ws.cell(row=r1, column=1).alignment = Alignment(
      horizontal="left", vertical="center", wrap_text=True
  )

  # 设置行高与列宽
  ws.row_dimensions[r1].height = 40
  ws.row_dimensions[r2].height = 25
  ws.row_dimensions[r3].height = 25

  ws.column_dimensions["A"].width = max(
      ws.column_dimensions["A"].width or 0, 16
  )
  ws.column_dimensions["B"].width = max(
      ws.column_dimensions["B"].width or 0, 14
  )
  for idx in range(len(damage_types)):
    col_letter = get_column_letter(3 + idx)
    current_w = ws.column_dimensions[col_letter].width or 0
    ws.column_dimensions[col_letter].width = max(current_w, 16)

  print(f"   ✅ [已生成汇总表] 要素: {element_name}")
  return r3 + 3  # 返回下一个表格的起始行（中间空 2 行）


def process_bridge_summary_to_new_file(
    input_file_path, output_file_path
):
  """主控函数：后台计算缓存后，直接读取真实扣分值并在新文件中纵向堆叠输出"""
  print(f"1. 正在准备处理源文件: {input_file_path}")

  # 🚀 第一步：后台调用 Excel 强制计算并保存缓存
  force_excel_recalculate_and_save(input_file_path)

  # 第二步：读取已缓存计算结果的原表（纯矩阵模式）
  df_raw = pd.read_excel(
      input_file_path, sheet_name="桥面系评分", header=None
  )
#   print(df_raw.iloc[:, :5])
  element_grouped_data = {}
  current_element = ""

  # 第三步：从第 10 行（索引 9，即明细数据第一行）开始精确遍历
  for idx in range(9, len(df_raw)):
    row = df_raw.iloc[idx]

    # 绝对列位置映射：
    # col 0: 要素, col 1: 损坏类型, col 2: 损害程度, col 3: 单项扣分DPhi
    col_element = str(row.iloc[0]).strip() if pd.notna(row.iloc[0]) else ""
    defect_type = str(row.iloc[1]).strip() if pd.notna(row.iloc[1]) else ""
    degree = str(row.iloc[2]).strip() if pd.notna(row.iloc[2]) else "无"
    raw_score = row.iloc[3] if len(row) > 3 else 0

    # 继承合并单元格的要素名称
    if col_element and col_element != "nan" and "要素" not in col_element:
      current_element = col_element

    # 过滤无效行或表尾说明行
    if (
        not current_element
        or current_element == "nan"
        or not defect_type
        or defect_type == "nan"
        or defect_type == "损坏类型"
    ):
      continue

    # 🚀 直接读取后台计算并缓存后的真实扣分值
    try:
      score = int(raw_score) if pd.notna(raw_score) else 0
    except (ValueError, TypeError):
      score = 0

    if current_element not in element_grouped_data:
      element_grouped_data[current_element] = {}

    element_grouped_data[current_element][defect_type] = {
        "degree": degree,
        "score": score,
    }

  print(
      f"   └─ 成功提取出 {len(element_grouped_data)}"
      " 个要素的数据，开始写入【新 Excel 文件】..."
  )

  # 第四步：创建全新的工作簿和工作表并写入纵向堆叠的汇总表
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = "桥面系评分汇总"
  ws.views.sheetView[0].showGridLines = True

  start_row = 1
  for element_name, element_data in element_grouped_data.items():
    start_row = draw_single_element_summary(
        ws, start_row, element_name, element_data
    )

  # 保存到新文件
  wb.save(output_file_path)
  print(f"\n🎉 结果已成功写入到新的 Excel 文件中: {output_file_path}")


if __name__ == "__main__":
  input_path = (
      r"G:\Python_lk\Python\BridgeDefectAnalyzer\input\城市桥梁评分计算.xlsx"
  )
  output_path = (
      r"G:\Python_lk\Python\BridgeDefectAnalyzer\output\桥面系各要素扣分汇总.xlsx"
  )

  process_bridge_summary_to_new_file(input_path, output_path)